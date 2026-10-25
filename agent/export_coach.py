#!/usr/bin/env python3
"""Export settimanale per il coach: "AAAAMMGG - Routine e storico carichi Hevy Luca.xlsx".

Un file, due fogli:
  (a) Allenamenti: una riga per serie (riscaldamenti inclusi e marcati), con data,
      seduta riconosciuta, esercizio, tipo di serie, kg, kg per mano (Stacco rumeno
      e Step Up: in Hevy è il totale dei due manubri), ripetizioni e "scarico"
      nelle settimane di scarico della scheda.
  (b) Routine: una riga per esercizio di ogni routine Hevy, con serie (riscaldamento
      e lavoro), range di ripetizioni e carichi.

Legge Hevy solo in GET, con paginazione (tramite morning.py). Il file contiene
dati reali: salvalo FUORI dal repository. Non stampa chiavi.

Uso:
    python agent/export_coach.py <cartella di uscita fuori dal repository>

Richiede openpyxl (non fa parte della libreria standard; morning.py non lo usa).
"""
import datetime as dt
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import morning  # noqa: E402

TYPE = {"warmup": "riscaldamento", "normal": "lavoro", "failure": "cedimento", "dropset": "dropset"}


def per_hand_from(scheda):
    """{titolo normalizzato: data da cui in Hevy si segna il totale dei due manubri}"""
    out = {}
    for sess in scheda["sessions"].values():
        for x in sess["exercises"]:
            if x.get("hevy_total_from"):
                for t in x["hevy_titles"]:
                    out[morning.norm(t)] = dt.date.fromisoformat(x["hevy_total_from"])
    return out


def scheda_ranges(scheda, tmap):
    """{codice esercizio Hevy: "10-12"} dai range di ripetizioni della scheda."""
    out = {}
    for sess in scheda["sessions"].values():
        for x in sess["exercises"]:
            rng = "%d" % x["rep_min"] if x["rep_min"] == x["rep_max"] else "%d-%d" % (x["rep_min"], x["rep_max"])
            if x.get("per_leg"):
                rng += " per gamba"
            for tid in morning.exercise_tids(x, tmap):
                out.setdefault(tid, rng)
    return out


def num_list(values):
    vals = ["corpo libero" if not v else morning.fmt_num(v) for v in values]
    return "/".join(vals) if len(set(vals)) > 1 else (vals[0] if vals else "")


def main(argv):
    if len(argv) != 2:
        sys.exit(__doc__)
    out_dir = os.path.abspath(argv[1])
    repo = os.path.dirname(HERE)
    if os.path.commonpath([out_dir, repo]) == repo:
        sys.exit("ERRORE: la cartella di uscita è dentro il repository; scegline una fuori (il file contiene dati reali).")
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        sys.exit("ERRORE: serve openpyxl (pip install openpyxl)")

    try:
        scheda = morning.load_scheda()
        workouts = sorted(morning.fetch_workouts(), key=lambda w: w["start_time"])
        routines = morning.fetch_routines()
    except morning.FetchError as e:
        sys.exit("ERRORE: %s" % e)

    classified = morning.classify(workouts, scheda)
    label = {id(w): l for w, l in classified}
    tmap = morning.template_map([(w, None) for w in workouts])
    it_title = {}
    for w in workouts:
        for e in w["exercises"]:
            it_title.setdefault(e["exercise_template_id"], e["title"])
    total_from = per_hand_from(scheda)
    ranges = scheda_ranges(scheda, tmap)
    today = morning.today_rome()

    F = "Arial"
    H = Font(name=F, bold=True, color="FFFFFF")
    HF = PatternFill("solid", fgColor="305496")
    DL = PatternFill("solid", fgColor="FCE4D6")
    WU = Font(name=F, color="808080")
    B = Font(name=F)
    thin = Side(style="thin", color="BFBFBF")
    BD = Border(top=thin, bottom=thin, left=thin, right=thin)
    WR = Alignment(wrap_text=True, vertical="top")

    def header(ws, row, cols, widths):
        for c, h in enumerate(cols, 1):
            cell = ws.cell(row, c, h)
            cell.font, cell.fill, cell.alignment, cell.border = H, HF, WR, BD
        for i, wdt in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = wdt

    def put(ws, r, vals, font=B, fill=None):
        for c, v in enumerate(vals, 1):
            cell = ws.cell(r, c, v)
            cell.font, cell.border, cell.alignment = font, BD, WR
            if fill:
                cell.fill = fill

    wb = Workbook()

    # (a) Allenamenti
    ws = wb.active
    ws.title = "Allenamenti"
    first = morning.to_rome(workouts[0]["start_time"]).strftime("%d/%m/%Y") if workouts else "-"
    last = morning.to_rome(workouts[-1]["start_time"]).strftime("%d/%m/%Y") if workouts else "-"
    ws["A1"] = "Allenamenti Hevy — %d allenamenti, %s - %s (scheda %s %s)" % (
        len(workouts), first, last, scheda["version"], scheda["version_id"])
    ws["A1"].font = Font(name=F, bold=True, size=13)
    ws["A2"] = ("Una riga per serie, riscaldamenti inclusi (in grigio). Orari Europe/Rome. 'Seduta' = riconoscimento dello script "
                "dagli esercizi. 'Scarico' = settimana di scarico della scheda (sedute escluse dalla progressione).")
    ws["A2"].font = Font(name=F, italic=True)
    header(ws, 4, ["Data", "Ora", "Seduta", "Scarico", "Titolo allenamento", "Esercizio", "Serie", "Tipo di serie",
                   "kg (come in Hevy)", "kg per mano", "Ripetizioni"],
           [11, 6, 10, 8, 36, 40, 6, 14, 10, 9, 10])
    r = 4
    for w in workouts:
        t = morning.to_rome(w["start_time"])
        lab = label.get(id(w)) or ("Core" if morning.norm(w["title"]).startswith("core") else "")
        deload = "scarico" if morning.in_deload(scheda, t.date()) else ""
        for e in w["exercises"]:
            tf = total_from.get(morning.norm(e["title"]))
            for j, st in enumerate(e["sets"], 1):
                r += 1
                kg = st.get("weight_kg")
                hand = round(kg / 2, 2) if (tf and kg is not None and t.date() >= tf) else None
                warm = st.get("type") == "warmup"
                put(ws, r, [t.date(), t.strftime("%H:%M"), lab, deload, w["title"].strip(), e["title"], j,
                            TYPE.get(st.get("type"), st.get("type")), kg, hand, st.get("reps")],
                    WU if warm else B, DL if deload else None)
                ws.cell(r, 1).number_format = "DD/MM/YYYY"
    ws.freeze_panes = "F5"
    ws.auto_filter.ref = "A4:K%d" % max(r, 5)

    # (b) Routine
    ws2 = wb.create_sheet("Routine")
    ws2["A1"] = "Routine Hevy al %s" % today.strftime("%d/%m/%Y")
    ws2["A1"].font = Font(name=F, bold=True, size=13)
    ws2["A2"] = ("Una riga per esercizio. Carichi come impostati in Hevy (Stacco rumeno e Step Up: totale dei due manubri, "
                 "più il carico per mano; Multipower: solo dischi). Range dalla routine se impostato, altrimenti dalla scheda.")
    ws2["A2"].font = Font(name=F, italic=True)
    header(ws2, 4, ["Routine", "N°", "Esercizio (Hevy italiano)", "Esercizio (routine, inglese)", "Serie riscaldamento",
                    "Serie di lavoro", "Range ripetizioni", "Carichi di lavoro (kg)", "Carichi per mano (kg)",
                    "Carichi riscaldamento (kg)", "Note"],
           [40, 5, 38, 36, 11, 9, 12, 20, 14, 18, 30])
    r = 4
    for rt in routines:
        for i, e in enumerate(rt.get("exercises", []), 1):
            tid = e.get("exercise_template_id")
            name = it_title.get(tid, "(mai registrato)")
            sets = e.get("sets", [])
            warm = [x for x in sets if x.get("type") == "warmup"]
            work = [x for x in sets if x.get("type") != "warmup"]
            rr = [x.get("rep_range") for x in work if x.get("rep_range")]
            if rr and rr[0].get("start") is not None:
                rng = "%s-%s" % (rr[0].get("start"), rr[0].get("end"))
            else:
                rng = ranges.get(tid, "")
            hand = ""
            if morning.norm(name) in total_from:
                hand = num_list([(x.get("weight_kg") or 0) / 2 for x in work])
            put(ws2, r + 1, [rt["title"].strip(), i, name, e.get("title"), len(warm), len(work), rng,
                             num_list([x.get("weight_kg") or 0 for x in work]), hand,
                             num_list([x.get("weight_kg") or 0 for x in warm]) if warm else "", e.get("notes") or ""])
            r += 1
    ws2.freeze_panes = "C5"
    ws2.auto_filter.ref = "A4:K%d" % max(r, 5)

    out = os.path.join(out_dir, "%s - Routine e storico carichi Hevy Luca.xlsx" % today.strftime("%Y%m%d"))
    wb.save(out)
    print(out)
    print("allenamenti: %d | routine: %d" % (len(workouts), len(routines)))
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
